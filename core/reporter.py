"""결과 저장: SQLite 이력, 결과 엑셀, ERP 업로드 양식, 협력사 통보서 초안, 원본 PDF 보관.

협력사에 자동 발송하지 않는다. 통보서는 파일 초안까지만 만든다.
"""

import logging
import re
import shutil
import sqlite3
from contextlib import closing
from datetime import datetime
from decimal import Decimal
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

from core.judge import fmt, limit_text
from core.models import CoaResult, SavedFiles, Status
from core.units import normalize_unit

log = logging.getLogger(__name__)

SCHEMA = """
CREATE TABLE IF NOT EXISTS documents (
    id              INTEGER PRIMARY KEY AUTOINCREMENT,
    processed_at    TEXT NOT NULL,
    source          TEXT NOT NULL,          -- mail / folder / upload
    mail_entry_id   TEXT,                   -- 메일 경로가 아니면 NULL
    attachment_name TEXT,                   -- 메일 첨부 원본 파일명
    sender          TEXT,
    received_at     TEXT,
    file_path       TEXT NOT NULL,          -- 보관된 원본 PDF
    supplier        TEXT,
    material_code   TEXT,
    product_name    TEXT,
    lot_no          TEXT,
    mfg_date        TEXT,
    overall_status  TEXT NOT NULL,
    reasons         TEXT,                   -- 문서 단위 검토 사유 ('; '로 연결)
    spec_version    TEXT,
    result_xlsx     TEXT,                   -- 결과 엑셀 경로
    notice_path     TEXT,                   -- 통보서 초안 경로 (FAIL만)
    UNIQUE (mail_entry_id, attachment_name) -- 같은 메일의 같은 첨부는 한 번만 (NULL은 중복 허용)
);
CREATE TABLE IF NOT EXISTS item_results (
    id               INTEGER PRIMARY KEY AUTOINCREMENT,
    document_id      INTEGER NOT NULL REFERENCES documents(id) ON DELETE CASCADE,
    item_std         TEXT,                  -- 판정 제외 항목은 NULL
    raw_name         TEXT,
    raw_value        TEXT,
    raw_unit         TEXT,
    normalized_value TEXT,                  -- Decimal 문자열 (규격 단위)
    spec_unit        TEXT,
    lower            TEXT,
    upper            TEXT,
    status           TEXT NOT NULL,         -- PASS / FAIL / REVIEW / EXCLUDED(Spec에 없는 항목)
    reason           TEXT
);
CREATE INDEX IF NOT EXISTS idx_documents_lot ON documents(lot_no);
CREATE INDEX IF NOT EXISTS idx_documents_entry ON documents(mail_entry_id);
CREATE INDEX IF NOT EXISTS idx_documents_processed ON documents(processed_at);
"""
ADDED_COLUMNS = {"result_xlsx": "TEXT", "notice_path": "TEXT"}  # STEP 8에서 추가: 기존 DB에도 열을 붙인다

EXCLUDED = "EXCLUDED"
FILL = {
    Status.FAIL: PatternFill("solid", fgColor="FFC7CE"),
    Status.REVIEW: PatternFill("solid", fgColor="FFEB9C"),
    Status.PASS: PatternFill("solid", fgColor="C6EFCE"),
}
HEADER_FILL = PatternFill("solid", fgColor="305496")
HEADER_FONT = Font(bold=True, color="FFFFFF")
STATUS_LABEL = {Status.PASS: "적합", Status.FAIL: "부적합", Status.REVIEW: "검토 필요"}
ERP_BASE_COLUMNS = ["처리일시", "협력사", "재질코드", "Lot No.", "판정"]


# ── 공통 ──────────────────────────────────────────────────────
def safe_part(text: str | None, default: str) -> str:
    """파일명 조각: 영숫자·한글·-만 남기고 나머지는 _ ('Alpha Chemical Co., Ltd.' → 'Alpha_Chemical_Co_Ltd')."""
    cleaned = re.sub(r"[^\w\-]+", "_", text or "").strip("_")
    return cleaned or default


def unique_path(path: Path) -> Path:
    """같은 이름 파일이 있으면 _2, _3 … 을 붙인다."""
    if not path.exists():
        return path
    n = 2
    while (candidate := path.with_name(f"{path.stem}_{n}{path.suffix}")).exists():
        n += 1
    return candidate


def _dec(value: Decimal | None) -> str | None:
    """Decimal을 DB 저장용 문자열로 (지수 표기 없음)."""
    return None if value is None else fmt(value)


def _as_datetime(value: Any) -> datetime | None:
    """datetime 또는 ISO 문자열을 datetime으로."""
    if value is None or isinstance(value, datetime):
        return value
    return datetime.fromisoformat(str(value))


# ── SQLite ────────────────────────────────────────────────────
def connect(db_path: Path) -> sqlite3.Connection:
    """쓰기용 연결. WAL 모드(쓰는 중에도 화면이 읽을 수 있음)로 열고, 테이블이 없으면 만든다."""
    db_path.parent.mkdir(parents=True, exist_ok=True)
    conn = sqlite3.connect(db_path, timeout=10)
    conn.execute("PRAGMA journal_mode = WAL")
    conn.execute("PRAGMA busy_timeout = 10000")
    conn.execute("PRAGMA foreign_keys = ON")
    conn.executescript(SCHEMA)
    have = {row[1] for row in conn.execute("PRAGMA table_info(documents)")}
    for column, kind in ADDED_COLUMNS.items():
        if column not in have:
            conn.execute(f"ALTER TABLE documents ADD COLUMN {column} {kind}")
    return conn


def connect_readonly(db_path: Path) -> sqlite3.Connection | None:
    """화면 조회용 읽기 전용 연결 (DB가 없으면 None). 조회할 때만 짧게 열고 바로 닫는다."""
    if not db_path.exists():
        return None
    conn = sqlite3.connect(f"{db_path.resolve().as_uri()}?mode=ro", uri=True, timeout=10)
    conn.execute("PRAGMA busy_timeout = 10000")
    conn.row_factory = sqlite3.Row
    return conn


def read_rows(db_path: Path, sql: str, params: tuple | list = ()) -> list[dict]:
    """읽기 전용 쿼리 결과를 dict 목록으로 (DB가 없으면 빈 목록)."""
    conn = connect_readonly(db_path)
    if conn is None:
        return []
    with closing(conn):
        return [dict(r) for r in conn.execute(sql, params).fetchall()]


def search_documents(db_path: Path, supplier: str | None = None, lot: str | None = None,
                     date_from: str | None = None, date_to: str | None = None,
                     statuses: list[str] | None = None, limit: int | None = None) -> list[dict]:
    """처리 이력 검색 (최근 처리 순). 날짜는 'YYYY-MM-DD', Lot은 부분 일치."""
    where, args = [], []
    if supplier:
        where.append("supplier = ?")
        args.append(supplier)
    if lot:
        where.append("lot_no LIKE ?")
        args.append(f"%{lot.strip()}%")
    if date_from:
        where.append("substr(processed_at, 1, 10) >= ?")
        args.append(date_from)
    if date_to:
        where.append("substr(processed_at, 1, 10) <= ?")
        args.append(date_to)
    if statuses:
        where.append(f"overall_status IN ({','.join('?' * len(statuses))})")
        args.extend(statuses)
    sql = "SELECT * FROM documents" + (" WHERE " + " AND ".join(where) if where else "")
    sql += " ORDER BY processed_at DESC, id DESC" + (f" LIMIT {int(limit)}" if limit else "")
    return read_rows(db_path, sql, args)


def get_document(db_path: Path, document_id: int) -> dict | None:
    """문서 1건."""
    rows = read_rows(db_path, "SELECT * FROM documents WHERE id = ?", (document_id,))
    return rows[0] if rows else None


def get_item_results(db_path: Path, document_ids: int | list[int]) -> list[dict]:
    """문서(들)의 항목별 판정 근거 (판정 제외 항목 포함, 저장 순서)."""
    ids = [document_ids] if isinstance(document_ids, int) else list(document_ids)
    if not ids:
        return []
    return read_rows(db_path, f"SELECT * FROM item_results WHERE document_id IN ({','.join('?' * len(ids))}) "
                              "ORDER BY document_id, id", ids)


def is_processed(db_path: Path, mail_entry_id: str, attachment_name: str | None = None) -> bool:
    """이 메일(EntryID)을 이미 처리했는지. attachment_name을 주면 그 첨부 단위로 확인한다."""
    if not mail_entry_id or not db_path.exists():
        return False
    sql = "SELECT 1 FROM documents WHERE mail_entry_id = ?"
    args: list = [mail_entry_id]
    if attachment_name is not None:
        sql += " AND attachment_name = ?"
        args.append(attachment_name)
    with closing(connect(db_path)) as conn:
        return conn.execute(sql + " LIMIT 1", args).fetchone() is not None


def is_file_recorded(db_path: Path, file_path: Path) -> bool:
    """이 PDF 경로가 이미 이력에 있는지 (폴더 감시 재시작 시 중복 처리 방지)."""
    if not db_path.exists():
        return False
    with closing(connect(db_path)) as conn:
        row = conn.execute("SELECT 1 FROM documents WHERE file_path = ? LIMIT 1", (str(Path(file_path).resolve()),))
        return row.fetchone() is not None


def save_to_db(db_path: Path, result: CoaResult, file_path: Path) -> int | None:
    """판정 결과와 항목별 근거를 저장하고 문서 ID를 반환한다. 같은 메일 첨부가 이미 있으면 None."""
    doc, meta = result.document, result.meta
    received = _as_datetime(meta.get("received_at"))
    row = (
        result.processed_at.isoformat(timespec="seconds"), meta.get("source", "upload"),
        meta.get("mail_entry_id"), meta.get("attachment_name"), meta.get("sender"),
        received.isoformat(timespec="seconds") if received else None, str(file_path),
        doc.supplier, doc.material_code, doc.product_name, doc.lot_no, doc.mfg_date,
        result.overall_status.value, "; ".join(result.reasons) or None, result.spec_version,
    )
    items = [
        (r.item_std, r.raw_name, r.raw_value, r.raw_unit, _dec(r.normalized_value), r.spec_unit,
         _dec(r.lower), _dec(r.upper), r.status.value, r.reason)
        for r in result.item_results
    ] + [
        (None, i.raw_name, i.raw_value, i.raw_unit, None, None, None, None, EXCLUDED, "Spec에 없는 항목 - 판정 제외")
        for i in result.extra_items
    ]
    with closing(connect(db_path)) as conn:
        try:
            with conn:
                cur = conn.execute(
                    "INSERT INTO documents (processed_at, source, mail_entry_id, attachment_name, sender, "
                    "received_at, file_path, supplier, material_code, product_name, lot_no, mfg_date, "
                    "overall_status, reasons, spec_version) VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)",
                    row,
                )
                doc_id = cur.lastrowid
                conn.executemany(
                    "INSERT INTO item_results (document_id, item_std, raw_name, raw_value, raw_unit, "
                    "normalized_value, spec_unit, lower, upper, status, reason) VALUES (?,?,?,?,?,?,?,?,?,?,?)",
                    [(doc_id, *i) for i in items],
                )
        except sqlite3.IntegrityError:
            log.info("이미 저장된 메일 첨부라 건너뜀: %s / %s", meta.get("mail_entry_id"), meta.get("attachment_name"))
            return None
    return doc_id


# ── 원본 PDF 보관 ─────────────────────────────────────────────
def archive_pdf(pdf_path: Path, inbox_dir: Path, result: CoaResult) -> Path:
    """원본 PDF를 inbox/{YYYYMMDD}/{수신일시}_{협력사}_{원본파일명}.pdf로 보관한다.

    이미 inbox 아래에 있으면 그대로 둔다(메일 경로는 저장 시 표준 이름을 쓴다).
    그 밖의 위치(samples, 업로드 임시 파일 등)는 원본을 지우지 않도록 복사한다.
    """
    pdf_path = Path(pdf_path).resolve()
    inbox_dir = inbox_dir.resolve()
    if pdf_path.is_relative_to(inbox_dir):
        return pdf_path
    when = _as_datetime(result.meta.get("received_at")) or result.processed_at
    original = result.meta.get("attachment_name") or pdf_path.name
    stem = re.sub(r'[\\/:*?"<>|]+', "_", Path(original).stem).strip() or "COA"
    name = f"{when:%Y%m%d_%H%M%S}_{safe_part(result.document.supplier, '협력사미상')}_{stem}.pdf"
    target = unique_path(inbox_dir / f"{when:%Y%m%d}" / name)
    target.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy2(pdf_path, target)
    return target


# ── 결과 엑셀 ─────────────────────────────────────────────────
def write_result_xlsx(results_dir: Path, result: CoaResult, pdf_path: Path) -> Path:
    """문서별 결과 엑셀: 상단 문서 정보, 아래 항목별 판정표(FAIL 빨강, REVIEW 노랑)."""
    doc = result.document
    name = (f"{result.processed_at:%Y%m%d_%H%M%S}_{safe_part(doc.supplier, '협력사미상')}_"
            f"{safe_part(doc.lot_no, 'Lot미상')}_결과.xlsx")
    path = unique_path(results_dir / name)

    wb = Workbook()
    ws = wb.active
    ws.title = "판정결과"
    status = result.overall_status
    ws.append(["COA 검증 결과"])
    ws["A1"].font = Font(bold=True, size=14)
    info = [
        ("문서 판정", f"{status.value} ({STATUS_LABEL[status]})"),
        ("Spec 버전", result.spec_version or "-"),
        ("협력사", doc.supplier or "-"),
        ("재질코드", doc.material_code or "-"),
        ("제품명", doc.product_name or "-"),
        ("Lot No.", doc.lot_no or "-"),
        ("제조일", doc.mfg_date or "-"),
        ("원본 파일", str(pdf_path)),
        ("처리일시", f"{result.processed_at:%Y-%m-%d %H:%M:%S}"),
        ("문서 검토 사유", "; ".join(result.reasons) or "-"),
    ]
    for label, value in info:
        ws.append([label, value])
        ws.cell(ws.max_row, 1).font = Font(bold=True)
    ws["B2"].fill = FILL[status]
    ws["B2"].font = Font(bold=True)
    ws.append([])

    header = ["항목", "원문 항목명", "원문 값", "원문 단위", "환산값", "규격 단위", "하한", "상한", "판정", "사유"]
    ws.append(header)
    header_row = ws.max_row
    for cell in ws[header_row]:
        cell.font, cell.fill = HEADER_FONT, HEADER_FILL
    for r in result.item_results:
        ws.append([r.item_std, r.raw_name, r.raw_value, r.raw_unit, r.normalized_value, r.spec_unit,
                   r.lower, r.upper, r.status.value, r.reason])
        if r.status is not Status.PASS:
            for cell in ws[ws.max_row]:
                cell.fill = FILL[r.status]
            ws.cell(ws.max_row, 9).font = Font(bold=True)
    for i in result.extra_items:
        ws.append(["(판정 제외)", i.raw_name, i.raw_value, i.raw_unit, None, None, None, None, EXCLUDED,
                   "Spec에 없는 항목 - 판정 제외"])

    for col, width in zip("ABCDEFGHIJ", (14, 40, 10, 10, 10, 10, 8, 8, 9, 48)):
        ws.column_dimensions[col].width = width
    for row in ws.iter_rows(min_row=header_row + 1):
        row[9].alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = ws.cell(header_row + 1, 1)

    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path


# ── ERP 업로드 양식 ────────────────────────────────────────────
def append_erp_row(output_dir: Path, result: CoaResult) -> Path:
    """일자별 ERP 업로드 양식(output/erp_upload_YYYYMMDD.xlsx)에 문서 1건을 1행으로 누적한다.

    항목별 측정값은 규격 단위로 환산한 값. 처음 나온 항목은 열을 오른쪽에 추가한다.
    """
    path = output_dir / f"erp_upload_{result.processed_at:%Y%m%d}.xlsx"
    if path.exists():
        wb = load_workbook(path)
        ws = wb.active
        header = [c.value for c in ws[1]]
    else:
        wb = Workbook()
        ws = wb.active
        ws.title = "ERP_UPLOAD"
        header = list(ERP_BASE_COLUMNS)

    values: dict[str, Any] = {
        "처리일시": f"{result.processed_at:%Y-%m-%d %H:%M:%S}",
        "협력사": result.document.supplier,
        "재질코드": result.document.material_code,
        "Lot No.": result.document.lot_no,
        "판정": result.overall_status.value,
    }
    for r in result.item_results:
        col = f"{r.item_std} ({r.spec_unit})"
        if col not in header:
            header.append(col)
        values[col] = r.normalized_value

    for i, name in enumerate(header, start=1):
        cell = ws.cell(1, i, name)
        cell.font, cell.fill = HEADER_FONT, HEADER_FILL
        ws.column_dimensions[cell.column_letter].width = max(12, len(name) + 4)
    ws.append([values.get(name) for name in header])
    ws.freeze_panes = "A2"

    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path


# ── 통보서 초안 ───────────────────────────────────────────────
def write_notice(notices_dir: Path, result: CoaResult) -> Path | None:
    """FAIL 문서만 협력사 부적합 통보서 초안(영문+국문 txt)을 만든다. 발송하지 않는다."""
    if result.overall_status is not Status.FAIL:
        return None
    doc = result.document
    fails = [r for r in result.item_results if r.status is Status.FAIL]
    supplier = doc.supplier or "[협력사명]"
    lot = doc.lot_no or "[Lot No.]"
    product = f"{doc.product_name or '-'} ({doc.material_code or '-'})"

    def line(r) -> str:
        reported = f"{r.raw_value} {r.raw_unit or ''}".strip()
        converted = (f" (= {fmt(r.normalized_value)} {r.spec_unit})"
                     if r.normalized_value is not None and normalize_unit(r.raw_unit) != r.spec_unit else "")
        return f"  - {r.item_std} ({r.raw_name}): {reported}{converted} / Spec {limit_text(r.lower, r.upper, r.spec_unit)}"

    items_text = "\n".join(line(r) for r in fails)
    text = f"""\
[부적합 통보서 초안] 자동 생성 문서입니다. 발송 전 담당자가 내용을 검토·수정하세요. (자동 발송되지 않음)
생성일시: {result.processed_at:%Y-%m-%d %H:%M}   원본 COA: {Path(result.document.source_path).name}

To   : Quality Assurance, {supplier}
Subj : Non-conformance Notice - {doc.material_code or 'Material'} Lot {lot}

Dear Quality Assurance Team,

Thank you for your continued support.
We reviewed the Certificate of Analysis for the lot below against our specification ({result.spec_version or '-'})
and found that the following item(s) do not conform.

  Product : {product}
  Lot No. : {lot}

{items_text}

We kindly ask you to provide (1) the root cause analysis, (2) corrective and preventive actions,
and (3) your proposal for the disposition of this lot (replacement or return) by [회신 기한].
Please let us know if you need any further information.

Best regards,
[담당자 이름]
Quality Control, [회사명]

------------------------------------------------------------------------
수신 : {supplier} 품질 담당자님
제목 : [부적합 통보] {doc.material_code or '원료'} Lot {lot} 시험성적서 규격 부적합 안내

안녕하십니까. 항상 협조해 주셔서 감사드립니다.
귀사에서 송부해 주신 아래 Lot의 시험성적서(COA)를 당사 규격({result.spec_version or '-'})과 대조한 결과,
다음 항목이 규격을 충족하지 않는 것으로 확인되어 알려 드립니다.

  제품명 : {product}
  Lot No.: {lot}

{items_text}

번거로우시겠지만 (1) 원인 조사 결과, (2) 재발 방지 대책, (3) 해당 Lot 처리 방안(대체 공급 또는 반품)을
[회신 기한]까지 회신해 주시면 감사하겠습니다.

[담당자 이름] 드림
품질관리팀, [회사명]
"""
    path = notices_dir / f"{safe_part(doc.lot_no, 'Lot미상')}_부적합통보서_초안.txt"
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(text, encoding="utf-8-sig")  # 메모장에서 한글이 깨지지 않게 BOM 포함
    return path


# ── 묶음 ──────────────────────────────────────────────────────
def save_all(result: CoaResult, pdf_path: Path, paths: dict[str, Path]) -> SavedFiles:
    """원본 보관 → DB → 결과 엑셀 → ERP 양식 → 통보서 순으로 저장한다.

    DB에 이미 있는 메일 첨부면 파일을 만들지 않고 document_id=None으로 돌려준다.
    엑셀이 열려 있는 등 파일 저장이 실패해도 나머지는 계속하고 errors에 남긴다.
    """
    saved = SavedFiles()
    output_dir = paths["output"]
    saved.pdf = archive_pdf(pdf_path, paths["inbox"], result)
    saved.document_id = save_to_db(paths["db"], result, saved.pdf)
    if saved.document_id is None:
        if saved.pdf != Path(pdf_path).resolve():
            saved.pdf.unlink(missing_ok=True)  # 방금 복사한 보관본은 되돌린다
        saved.pdf = None
        saved.errors.append("이미 처리한 메일 첨부라 저장하지 않음")
        return saved

    steps = (
        ("result_xlsx", "결과 엑셀", lambda: write_result_xlsx(output_dir / "results", result, saved.pdf)),
        ("erp_xlsx", "ERP 업로드 양식", lambda: append_erp_row(output_dir, result)),
        ("notice", "통보서 초안", lambda: write_notice(output_dir / "notices", result)),
    )
    for field, label, make in steps:
        try:
            setattr(saved, field, make())
        except PermissionError as exc:
            msg = f"{label} 저장 실패 - 파일이 다른 프로그램(엑셀 등)에서 열려 있는지 확인: {exc.filename}"
            log.warning(msg)
            saved.errors.append(msg)
    with closing(connect(paths["db"])) as conn, conn:
        conn.execute("UPDATE documents SET result_xlsx = ?, notice_path = ? WHERE id = ?",
                     (str(saved.result_xlsx) if saved.result_xlsx else None,
                      str(saved.notice) if saved.notice else None, saved.document_id))
    return saved
