"""spec_master.xlsx 로딩 (품목별 항목·하한·상한·단위·별칭·협력사 도메인).

숫자는 Decimal(str(값))으로 변환해 엑셀 float의 이진 오차가 판정에 끼지 않게 한다.
"""

import re
from dataclasses import dataclass
from decimal import Decimal, InvalidOperation
from pathlib import Path

from openpyxl import load_workbook

from core.config import load_config
from core.units import normalize_unit

SHEET_COLUMNS = {
    "materials": ["material_code", "material_name", "name_aliases"],
    "spec": ["material_code", "item_std", "aliases", "lower", "upper", "unit", "spec_version"],
    "suppliers": ["supplier_name", "email_domain", "name_aliases"],
}


@dataclass(frozen=True)
class Material:
    """재질(원료) 1건: 코드, 정식 이름, 이름 별칭."""

    code: str
    name: str
    aliases: tuple[str, ...]


@dataclass(frozen=True)
class SpecItem:
    """재질별 규격 항목 1건. lower/upper가 None이면 그쪽 한계 없음."""

    material_code: str
    item_std: str
    aliases: tuple[str, ...]
    lower: Decimal | None
    upper: Decimal | None
    unit: str
    spec_version: str


@dataclass(frozen=True)
class Supplier:
    """협력사 1건: 정식 이름, 메일 도메인, 이름 별칭."""

    name: str
    email_domain: str
    aliases: tuple[str, ...]


@dataclass(frozen=True)
class SpecMaster:
    """규격 마스터 전체와 조회 함수."""

    materials: tuple[Material, ...]
    specs: tuple[SpecItem, ...]
    suppliers: tuple[Supplier, ...]

    def specs_for(self, material_code: str) -> list[SpecItem]:
        """재질 코드의 규격 항목 목록."""
        return [s for s in self.specs if s.material_code == material_code]

    def item_aliases(self, material_code: str | None = None) -> dict[str, str]:
        """항목 별칭 → item_std 사전 (item_std 자신 포함). material_code를 주면 그 재질만."""
        specs = self.specs_for(material_code) if material_code else self.specs
        return {alias: s.item_std for s in specs for alias in (s.item_std, *s.aliases)}

    def material_aliases(self) -> dict[str, str]:
        """재질 이름·별칭 → material_code 사전 (코드·정식 이름 포함)."""
        return {alias: m.code for m in self.materials for alias in (m.code, m.name, *m.aliases)}

    def spec_version(self, material_code: str) -> str | None:
        """재질의 Spec 버전. 항목별 버전이 다르면 쉼표로 이어 붙인다."""
        versions = sorted({s.spec_version for s in self.specs_for(material_code)})
        return ", ".join(versions) if versions else None

    def supplier_by_domain(self, email_or_domain: str) -> Supplier | None:
        """메일 주소나 도메인으로 협력사를 찾는다 (하위 도메인 포함, 대소문자 무시)."""
        domain = email_or_domain.rsplit("@", 1)[-1].strip().lower()
        for s in self.suppliers:
            d = s.email_domain.lower()
            if domain == d or domain.endswith("." + d):
                return s
        return None

    @property
    def supplier_domains(self) -> list[str]:
        """등록된 협력사 메일 도메인 목록."""
        return [s.email_domain.lower() for s in self.suppliers]


def _text(value) -> str:
    """셀 값을 앞뒤 공백 없는 문자열로 (빈 셀은 "")."""
    return "" if value is None else str(value).strip()


def _split_aliases(value) -> tuple[str, ...]:
    """'a, b; c' 형태의 별칭 칸을 튜플로 나눈다."""
    return tuple(a.strip() for a in re.split(r"[,;\n]", _text(value)) if a.strip())


def to_decimal(value, where: str) -> Decimal | None:
    """엑셀 숫자를 Decimal(str(값))으로 바꾼다. 빈 칸은 None, 숫자가 아니면 ValueError."""
    text = _text(value)
    if not text:
        return None
    try:
        return Decimal(text)
    except InvalidOperation:
        raise ValueError(f"규격 마스터 {where}: 숫자가 아닙니다: {text!r}") from None


def _read_rows(wb, sheet: str) -> list[dict]:
    """시트의 헤더를 확인하고 행을 {열 이름: 값} 목록으로 읽는다 (빈 행 제외)."""
    if sheet not in wb.sheetnames:
        raise ValueError(f"규격 마스터에 '{sheet}' 시트가 없습니다.")
    rows = wb[sheet].iter_rows(values_only=True)
    header = [_text(h) for h in next(rows, ())]
    missing = [c for c in SHEET_COLUMNS[sheet] if c not in header]
    if missing:
        raise ValueError(f"규격 마스터 '{sheet}' 시트에 열이 없습니다: {', '.join(missing)}")
    result = []
    for row in rows:
        if all(_text(v) == "" for v in row):
            continue
        result.append(dict(zip(header, row)))
    return result


def load_spec_master(path: Path | str | None = None) -> SpecMaster:
    """spec_master.xlsx를 읽어 SpecMaster로 반환한다. 경로를 안 주면 config.toml의 paths.spec_master."""
    path = Path(path) if path else load_config()["paths"]["spec_master"]
    if not path.is_file():
        raise FileNotFoundError(f"규격 마스터 파일을 찾을 수 없습니다: {path}")
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        materials = tuple(
            Material(_text(r["material_code"]), _text(r["material_name"]), _split_aliases(r["name_aliases"]))
            for r in _read_rows(wb, "materials")
        )
        specs = []
        for n, r in enumerate(_read_rows(wb, "spec"), start=2):
            where = f"spec 시트 {n}행"
            unit = normalize_unit(_text(r["unit"]))
            if unit is None:
                raise ValueError(f"규격 마스터 {where}: 모르는 단위입니다: {_text(r['unit'])!r}")
            specs.append(SpecItem(
                material_code=_text(r["material_code"]),
                item_std=_text(r["item_std"]),
                aliases=_split_aliases(r["aliases"]),
                lower=to_decimal(r["lower"], f"{where} lower"),
                upper=to_decimal(r["upper"], f"{where} upper"),
                unit=unit,
                spec_version=_text(r["spec_version"]),
            ))
        suppliers = tuple(
            Supplier(_text(r["supplier_name"]), _text(r["email_domain"]), _split_aliases(r["name_aliases"]))
            for r in _read_rows(wb, "suppliers")
        )
    finally:
        wb.close()
    return SpecMaster(materials, tuple(specs), suppliers)
