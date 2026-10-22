"""scripts/make_samples.py 생성물 테스트: 임시 폴더에 다시 만들어 내용을 확인."""

import json

import pdfplumber
from openpyxl import load_workbook

from scripts import make_samples


def test_build_all(tmp_path):
    """PDF 5장·expected.json·spec_master.xlsx가 생기고 PDF 텍스트에서 Lot과 항목값이 읽힌다."""
    samples_dir, spec_path = tmp_path / "samples", tmp_path / "spec_master.xlsx"
    pdfs = make_samples.build_all(samples_dir, spec_path)

    assert len(pdfs) == 5 and all(p.is_file() for p in pdfs)
    assert make_samples.verify_pdfs(samples_dir)

    wb = load_workbook(spec_path)
    assert wb.sheetnames == ["materials", "spec", "suppliers"]
    assert [c.value for c in wb["spec"][1]] == make_samples.SPEC_HEADER
    assert wb["spec"].max_row == 5


def test_expected_overall(tmp_path):
    """정답지의 문서 판정이 파일명 접미사(pass/fail/review)와 일치한다."""
    make_samples.build_all(tmp_path / "samples", tmp_path / "spec.xlsx")
    data = json.loads((tmp_path / "samples" / "expected.json").read_text(encoding="utf-8"))
    for name, sample in data["samples"].items():
        assert sample["overall"] == name.rsplit("_", 1)[1].removesuffix(".pdf").upper()
        assert set(sample["items"]) == {"Purity", "Moisture", "Iron", "Color"}


def test_alpha_table_extracted(tmp_path):
    """Alpha 양식은 pdfplumber 표 추출로 헤더와 4개 항목 행이 나온다."""
    make_samples.build_all(tmp_path / "samples", tmp_path / "spec.xlsx")
    with pdfplumber.open(tmp_path / "samples" / "COA_Alpha_A2402_fail.pdf") as pdf:
        table = pdf.pages[0].extract_tables()[0]
    assert table[0] == ["Test Item", "Specification", "Result", "Unit"]
    assert ["Iron (Fe)", "Max. 0.005", "0.008", "%"] in table
