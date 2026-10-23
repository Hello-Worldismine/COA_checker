"""core.reporter 테스트: 임시 폴더·임시 DB로 샘플 5장을 처리해 저장 결과 확인."""

import json
import sqlite3
from datetime import datetime

import pytest
from openpyxl import load_workbook

from core import reporter
from core.config import load_config
from core.pipeline import process_pdf
from core.reporter import is_processed

SAMPLES_DIR = load_config()["paths"]["samples"]
EXPECTED = json.loads((SAMPLES_DIR / "expected.json").read_text(encoding="utf-8"))["samples"]
RECEIVED = datetime(2026, 10, 23, 9, 30, 15)


@pytest.fixture
def paths(tmp_path):
    """임시 inbox·output·DB 경로."""
    return {"inbox": tmp_path / "inbox", "output": tmp_path / "output", "db": tmp_path / "output" / "test.db"}


def _mail_meta(name, entry_id="ENTRY-1"):
    """메일 1통에 첨부된 PDF의 meta."""
    return {"source": "mail", "mail_entry_id": entry_id, "attachment_name": name,
            "sender": "jaemin9905@naver.com", "received_at": RECEIVED}


def _process_all(paths, spec_master, entry_id="ENTRY-1"):
    """샘플 5장을 같은 메일의 첨부로 처리한다."""
    return {name: process_pdf(SAMPLES_DIR / name, _mail_meta(name, entry_id), spec=spec_master, paths=paths)
            for name in sorted(EXPECTED)}


def _count(db, table, where="1=1"):
    with sqlite3.connect(db) as conn:
        return conn.execute(f"SELECT COUNT(*) FROM {table} WHERE {where}").fetchone()[0]


def test_five_samples_saved(paths, spec_master):
    """DB 문서 5행·항목 20행, 결과 엑셀 5개, ERP 양식 1개(5행), FAIL 2건에만 통보서."""
    outs = _process_all(paths, spec_master)
    db = paths["db"]

    assert all(not o.skipped and o.saved.errors == [] for o in outs.values())
    assert _count(db, "documents") == 5
    assert _count(db, "item_results") == 20
    assert _count(db, "documents", "overall_status='FAIL'") == 2

    results = sorted((paths["output"] / "results").glob("*_결과.xlsx"))
    assert len(results) == 5
    erp = list(paths["output"].glob("erp_upload_*.xlsx"))
    assert len(erp) == 1
    ws = load_workbook(erp[0]).active
    assert ws.max_row == 6  # 머리글 + 5건
    assert [c.value for c in ws[1]][:5] == ["처리일시", "협력사", "재질코드", "Lot No.", "판정"]

    notices = sorted(p.name for p in (paths["output"] / "notices").glob("*.txt"))
    assert notices == ["A2402_부적합통보서_초안.txt", "N2418_부적합통보서_초안.txt"]
    for name, out in outs.items():
        assert (out.saved.notice is not None) == (EXPECTED[name]["overall"] == "FAIL")


def test_duplicate_entry_id_not_saved_twice(paths, spec_master):
    """같은 mail_entry_id(같은 첨부)를 다시 처리하면 건너뛰고 아무것도 늘지 않는다."""
    _process_all(paths, spec_master)
    assert is_processed(paths["db"], "ENTRY-1")
    assert is_processed(paths["db"], "ENTRY-1", "COA_Alpha_A2401_pass.pdf")
    assert not is_processed(paths["db"], "ENTRY-2")

    again = _process_all(paths, spec_master)
    assert all(o.skipped and o.result is None for o in again.values())
    assert _count(paths["db"], "documents") == 5
    assert len(list((paths["output"] / "results").glob("*.xlsx"))) == 5
    assert len(list((paths["inbox"]).rglob("*.pdf"))) == 5


def test_save_to_db_rejects_duplicate_directly(paths, spec_master):
    """사전 확인을 건너뛰어도 DB UNIQUE 제약으로 중복이 막히고 복사본도 남지 않는다."""
    name = "COA_Alpha_A2401_pass.pdf"
    first = process_pdf(SAMPLES_DIR / name, _mail_meta(name), spec=spec_master, paths=paths)
    second = reporter.save_all(first.result, SAMPLES_DIR / name, paths)
    assert second.document_id is None and second.pdf is None
    assert _count(paths["db"], "documents") == 1
    assert len(list(paths["inbox"].rglob("*.pdf"))) == 1


def test_upload_without_entry_id_saved_each_time(paths, spec_master):
    """EntryID가 없는 업로드·폴더 경로는 매번 새 이력으로 저장한다."""
    pdf = SAMPLES_DIR / "COA_Nippon_N2418_fail.pdf"
    for _ in range(2):
        out = process_pdf(pdf, {"source": "upload"}, spec=spec_master, paths=paths)
        assert not out.skipped and out.saved.document_id
    assert _count(paths["db"], "documents", "source='upload'") == 2
    assert len(list((paths["output"] / "results").glob("*.xlsx"))) == 2  # 같은 초에 처리돼도 _2로 따로


def test_db_keeps_judgment_evidence(paths, spec_master):
    """항목별 원문 값·단위·환산값(TEXT)·규격·판정·사유·Spec 버전을 남긴다."""
    name = "COA_Nippon_N2418_fail.pdf"
    process_pdf(SAMPLES_DIR / name, _mail_meta(name), spec=spec_master, paths=paths)
    with sqlite3.connect(paths["db"]) as conn:
        doc = conn.execute("SELECT source, sender, received_at, lot_no, supplier, overall_status, spec_version "
                           "FROM documents").fetchone()
        item = conn.execute("SELECT raw_name, raw_value, raw_unit, normalized_value, spec_unit, lower, upper, "
                            "status, reason FROM item_results WHERE item_std='Moisture'").fetchone()
    assert doc == ("mail", "jaemin9905@naver.com", "2026-10-23T09:30:15", "N2418",
                   "Nippon Solvent Co., Ltd.", "FAIL", "v1.0")
    assert item == ("Water", "50", "ppm", "0.005", "%", None, "0.003", "FAIL",
                    "50 ppm → 0.005 % 환산, 상한 0.003 % 초과")


def test_pdf_archived_with_standard_name(paths, spec_master):
    """원본은 inbox/YYYYMMDD/{수신일시}_{협력사}_{원본명}.pdf로 복사 보관 (samples 원본은 그대로)."""
    name = "COA_Alpha_A2402_fail.pdf"
    out = process_pdf(SAMPLES_DIR / name, _mail_meta(name), spec=spec_master, paths=paths)
    assert out.saved.pdf == paths["inbox"].resolve() / "20261023" / \
        "20261023_093015_Alpha_Chemical_Co_Ltd_COA_Alpha_A2402_fail.pdf"
    assert out.saved.pdf.is_file() and (SAMPLES_DIR / name).is_file()


def test_pdf_already_in_inbox_not_moved(paths, spec_master):
    """이미 inbox 아래에 있는 PDF는 옮기거나 복사하지 않는다."""
    target = paths["inbox"] / "20261023" / "already.pdf"
    target.parent.mkdir(parents=True)
    target.write_bytes((SAMPLES_DIR / "COA_Alpha_A2401_pass.pdf").read_bytes())
    out = process_pdf(target, {"source": "folder"}, spec=spec_master, paths=paths)
    assert out.saved.pdf == target.resolve()
    assert len(list(paths["inbox"].rglob("*.pdf"))) == 1


def test_result_xlsx_colors(paths, spec_master):
    """결과 엑셀: 상단 문서 판정·Spec 버전, FAIL 행 빨강, REVIEW 행 노랑."""
    fail = process_pdf(SAMPLES_DIR / "COA_Alpha_A2402_fail.pdf", {"source": "upload"}, spec=spec_master, paths=paths)
    ws = load_workbook(fail.saved.result_xlsx).active
    assert ws["B2"].value == "FAIL (부적합)" and ws["B3"].value == "v1.0"
    rows = {r[0].value: r for r in ws.iter_rows() if r[8].value in ("PASS", "FAIL", "REVIEW")}
    assert rows["Iron"][8].fill.fgColor.rgb.endswith("FFC7CE")
    assert rows["Purity"][8].fill.fgColor.rgb in (None, "00000000")

    review = process_pdf(SAMPLES_DIR / "COA_Nippon_N2419_review.pdf", {"source": "upload"}, spec=spec_master,
                         paths=paths)
    ws = load_workbook(review.saved.result_xlsx).active
    iron = next(r for r in ws.iter_rows() if r[0].value == "Iron")
    assert iron[8].value == "REVIEW" and iron[8].fill.fgColor.rgb.endswith("FFEB9C")


def test_notice_content(paths, spec_master):
    """통보서 초안: 협력사·Lot·부적합 항목·측정값·규격, 영문+국문, 자동 발송 안 함 문구."""
    out = process_pdf(SAMPLES_DIR / "COA_Nippon_N2418_fail.pdf", {"source": "upload"}, spec=spec_master,
                      paths=paths)
    text = out.saved.notice.read_text(encoding="utf-8-sig")
    assert "Nippon Solvent Co., Ltd." in text and "Lot N2418" in text
    assert "Moisture (Water): 50 ppm (= 0.005 %) / Spec ≤ 0.003 %" in text
    assert "Dear Quality Assurance Team" in text and "안녕하십니까" in text
    assert "자동 발송되지 않음" in text
    assert "Fe" not in text.split("Lot No.")[1].split("We kindly")[0]  # 적합 항목은 나열하지 않음


def test_erp_values_in_spec_unit(paths, spec_master):
    """ERP 양식 측정값은 규격 단위로 환산한 값, REVIEW 항목(N.D.)은 빈칸."""
    for name in ("COA_Nippon_N2418_fail.pdf", "COA_Nippon_N2419_review.pdf"):
        out = process_pdf(SAMPLES_DIR / name, {"source": "upload"}, spec=spec_master, paths=paths)
    ws = load_workbook(out.saved.erp_xlsx).active
    header = [c.value for c in ws[1]]
    rows = {r[3]: dict(zip(header, r)) for r in ws.iter_rows(min_row=2, values_only=True)}
    assert float(rows["N2418"]["Moisture (%)"]) == pytest.approx(0.005)
    assert rows["N2418"]["판정"] == "FAIL"
    assert rows["N2419"]["Iron (%)"] is None


def test_locked_excel_does_not_stop_saving(paths, spec_master, monkeypatch):
    """엑셀 파일이 열려 있어 저장이 막혀도 DB 저장과 나머지 파일은 계속된다."""
    def locked(*args, **kwargs):
        raise PermissionError(13, "Permission denied", "erp_upload.xlsx")

    monkeypatch.setattr(reporter, "append_erp_row", locked)
    out = process_pdf(SAMPLES_DIR / "COA_Alpha_A2402_fail.pdf", {"source": "upload"}, spec=spec_master, paths=paths)
    assert out.saved.document_id and out.saved.result_xlsx and out.saved.notice
    assert out.saved.erp_xlsx is None
    assert any("ERP 업로드 양식 저장 실패" in e for e in out.saved.errors)


def test_save_false_writes_nothing(paths, spec_master):
    """save=False는 판정만 하고 아무 파일도 만들지 않는다."""
    out = process_pdf(SAMPLES_DIR / "COA_Alpha_A2401_pass.pdf", spec=spec_master, save=False, paths=paths)
    assert out.result.overall_status == "PASS" and out.saved is None
    assert not paths["db"].exists() and not paths["output"].exists()
