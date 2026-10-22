"""개발용 샘플 COA PDF와 spec_master.xlsx 생성 (reportlab, openpyxl).

실행: .venv\\Scripts\\python scripts\\make_samples.py   (몇 번을 실행해도 같은 내용으로 다시 만든다)
"""

import json
import sys
from dataclasses import dataclass
from pathlib import Path

PROJECT_ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(PROJECT_ROOT))
sys.stdout.reconfigure(encoding="utf-8")  # 파일로 저장해도 한글이 깨지지 않게

import pdfplumber  # noqa: E402
from openpyxl import Workbook  # noqa: E402
from openpyxl.styles import Font, PatternFill  # noqa: E402
from reportlab import rl_config  # noqa: E402
from reportlab.lib import colors  # noqa: E402
from reportlab.lib.pagesizes import A4  # noqa: E402
from reportlab.lib.styles import getSampleStyleSheet  # noqa: E402
from reportlab.lib.units import mm  # noqa: E402
from reportlab.pdfgen import canvas  # noqa: E402
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle  # noqa: E402

from core.config import load_config  # noqa: E402

rl_config.invariant = 1  # 생성 시각·문서 ID를 고정해 다시 만들어도 같은 PDF가 나오게

SPEC_VERSION = "v1.0"
MATERIAL_CODE = "NMP-EL"

# ── 규격 마스터 ──────────────────────────────────────────────
MATERIALS_HEADER = ["material_code", "material_name", "name_aliases"]
MATERIALS = [
    [MATERIAL_CODE, "N-Methyl-2-pyrrolidone Electronic Grade",
     "NMP, N-Methyl-2-pyrrolidone, 1-Methyl-2-pyrrolidinone"],
]

SPEC_HEADER = ["material_code", "item_std", "aliases", "lower", "upper", "unit", "spec_version"]
SPEC = [
    [MATERIAL_CODE, "Purity", "Purity, Assay, Purity (GC), 순도", 99.90, None, "%", SPEC_VERSION],
    [MATERIAL_CODE, "Moisture", "Moisture, Water, Water content, H2O, KF, 수분", None, 0.003, "%", SPEC_VERSION],
    [MATERIAL_CODE, "Iron", "Iron, Fe, Iron (Fe), 철", None, 0.005, "%", SPEC_VERSION],
    [MATERIAL_CODE, "Color", "Color, Colour, APHA, 색도", None, 10, "APHA", SPEC_VERSION],
]

SUPPLIERS_HEADER = ["supplier_name", "email_domain", "name_aliases"]
SUPPLIERS = [
    ["Alpha Chemical Co., Ltd.", "alpha-chem.example", "Alpha Chemical"],
    ["Nippon Solvent Co., Ltd.", "nippon-solv.example", "Nippon Solvent"],
]


# ── 샘플 COA 정의 (= 정답지) ───────────────────────────────────
@dataclass
class Item:
    """COA 한 줄: 사내 표준 항목, COA 원문 항목명·값·단위, 기대 판정, 판정 사유."""

    std: str
    name: str
    value: str
    unit: str
    status: str
    note: str = ""


@dataclass
class Sample:
    """샘플 COA 1장의 내용과 기대 판정."""

    file: str
    form: str  # "alpha"(표 양식) / "nippon"(줄 나열 양식)
    supplier: str
    lot: str
    mfg_date: str
    items: list[Item]

    @property
    def overall(self) -> str:
        """문서 판정: FAIL이 하나라도 있으면 FAIL, 아니면 REVIEW가 있으면 REVIEW, 모두 PASS면 PASS."""
        statuses = {i.status for i in self.items}
        return "FAIL" if "FAIL" in statuses else "REVIEW" if "REVIEW" in statuses else "PASS"


ALPHA = "Alpha Chemical Co., Ltd."
NIPPON = "Nippon Solvent Co., Ltd."

ALPHA_SPEC_TEXT = {"Purity": "Min. 99.90", "Moisture": "Max. 0.003", "Iron": "Max. 0.005", "Color": "Max. 10"}


def _alpha_items(iron: str = "0.003") -> list[Item]:
    """Alpha 양식 기본 항목(모두 적합). iron 값만 바꿀 수 있다."""
    iron_fail = iron == "0.008"
    return [
        Item("Purity", "Purity", "99.95", "%", "PASS"),
        Item("Moisture", "Moisture", "0.002", "%", "PASS"),
        Item("Iron", "Iron (Fe)", iron, "%", "FAIL" if iron_fail else "PASS",
             "0.008 % > 상한 0.005 %" if iron_fail else ""),
        Item("Color", "Color", "5", "APHA", "PASS"),
    ]


def _nippon_items(water: tuple[str, str] = ("20", "ppm"), fe: tuple[str, str] = ("10", "ppm")) -> list[Item]:
    """Nippon 양식 기본 항목(모두 적합, ppm 환산 필요). Water·Fe 값만 바꿀 수 있다."""
    water_status, water_note = {
        "20": ("PASS", "20 ppm = 0.002 % ≤ 상한 0.003 %"),
        "50": ("FAIL", "50 ppm = 0.005 % > 상한 0.003 %"),
        "<10": ("PASS", "<10 ppm: 10 ppm = 0.001 % ≤ 상한 0.003 %"),
    }[water[0]]
    fe_status, fe_note = {
        "10": ("PASS", "10 ppm = 0.001 % ≤ 상한 0.005 %"),
        "N.D.": ("REVIEW", "N.D. 표기: 검출한계 불명"),
    }[fe[0]]
    return [
        Item("Purity", "Assay", "99.93", "%", "PASS"),
        Item("Moisture", "Water", water[0], water[1], water_status, water_note),
        Item("Iron", "Fe", fe[0], fe[1], fe_status, fe_note),
        Item("Color", "Colour", "8", "APHA", "PASS"),
    ]


SAMPLES = [
    Sample("COA_Alpha_A2401_pass.pdf", "alpha", ALPHA, "A2401", "2026-10-08", _alpha_items()),
    Sample("COA_Nippon_N2417_pass.pdf", "nippon", NIPPON, "N2417", "2026-10-12", _nippon_items()),
    Sample("COA_Alpha_A2402_fail.pdf", "alpha", ALPHA, "A2402", "2026-10-15", _alpha_items(iron="0.008")),
    Sample("COA_Nippon_N2418_fail.pdf", "nippon", NIPPON, "N2418", "2026-10-16", _nippon_items(water=("50", "ppm"))),
    Sample("COA_Nippon_N2419_review.pdf", "nippon", NIPPON, "N2419", "2026-10-19",
           _nippon_items(water=("<10", "ppm"), fe=("N.D.", ""))),
]


# ── 생성 함수 ─────────────────────────────────────────────────
def write_spec_master(path: Path) -> None:
    """spec_master.xlsx를 materials·spec·suppliers 시트로 만든다."""
    wb = Workbook()
    wb.remove(wb.active)
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="305496")
    for title, header, rows in (
        ("materials", MATERIALS_HEADER, MATERIALS),
        ("spec", SPEC_HEADER, SPEC),
        ("suppliers", SUPPLIERS_HEADER, SUPPLIERS),
    ):
        ws = wb.create_sheet(title)
        ws.append(header)
        for cell in ws[1]:
            cell.font, cell.fill = header_font, header_fill
        for row in rows:
            ws.append(row)
        for col in ws.columns:
            width = max(len(str(c.value)) if c.value is not None else 0 for c in col)
            ws.column_dimensions[col[0].column_letter].width = min(width + 2, 60)
        ws.freeze_panes = "A2"
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


def build_alpha_pdf(path: Path, s: Sample) -> None:
    """Alpha 양식: 상단 회사·제품·Lot 정보, 아래 시험 결과 표(Test Item | Specification | Result | Unit)."""
    styles = getSampleStyleSheet()
    doc = SimpleDocTemplate(str(path), pagesize=A4, title=f"COA {s.lot}", author=ALPHA,
                            leftMargin=20 * mm, rightMargin=20 * mm, topMargin=20 * mm)
    info = [
        "<b>Product Name:</b> N-Methyl-2-pyrrolidone (NMP) Electronic Grade",
        f"<b>Lot No.:</b> {s.lot}",
        f"<b>Mfg. Date:</b> {s.mfg_date}",
        "<b>Quantity:</b> 20 drums x 200 kg",
    ]
    rows = [["Test Item", "Specification", "Result", "Unit"]]
    rows += [[i.name, ALPHA_SPEC_TEXT[i.std], i.value, i.unit] for i in s.items]
    table = Table(rows, colWidths=[55 * mm, 45 * mm, 35 * mm, 30 * mm])
    table.setStyle(TableStyle([
        ("GRID", (0, 0), (-1, -1), 0.6, colors.black),
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#D9E1F2")),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("ALIGN", (1, 0), (-1, -1), "CENTER"),
        ("TOPPADDING", (0, 0), (-1, -1), 5),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
    ]))
    story = [
        Paragraph(ALPHA, styles["Title"]),
        Paragraph("123 Industrial Road, Ulsan, Korea  |  qc@alpha-chem.example", styles["Normal"]),
        Spacer(1, 8 * mm),
        Paragraph("CERTIFICATE OF ANALYSIS", styles["Heading2"]),
        *[Paragraph(line, styles["Normal"]) for line in info],
        Spacer(1, 6 * mm),
        table,
        Spacer(1, 12 * mm),
        Paragraph("Approved by: QC Manager, Alpha Chemical Co., Ltd.", styles["Normal"]),
    ]
    doc.build(story)


def build_nippon_pdf(path: Path, s: Sample) -> None:
    """Nippon 양식: 표 없이 'Assay ........ 99.93 %'처럼 항목을 줄로 나열, 단위는 값 뒤에 붙는다."""
    c = canvas.Canvas(str(path), pagesize=A4)
    c.setTitle(f"Certificate of Analysis {s.lot}")
    c.setAuthor(NIPPON)
    x, y = 25 * mm, A4[1] - 30 * mm

    def line(text: str, font: str = "Courier", size: int = 11, gap: float = 6.5) -> None:
        nonlocal y
        c.setFont(font, size)
        c.drawString(x, y, text)
        y -= gap * mm

    line(NIPPON, "Helvetica-Bold", 16, 7)
    line("Quality Assurance Dept.  qa@nippon-solv.example", "Helvetica", 9, 12)
    line("CERTIFICATE OF ANALYSIS", "Helvetica-Bold", 13, 10)
    line("Product Name : 1-Methyl-2-pyrrolidinone (NMP), EL Grade")
    line(f"Batch No.    : {s.lot}")
    line(f"Date of Mfg. : {s.mfg_date}", gap=12)
    line("Test Results", "Helvetica-Bold", 11, 8)
    for i in s.items:
        result = f"{i.value} {i.unit}".strip()
        line(f"{i.name} {'.' * (28 - len(i.name))} {result}")
    y -= 10 * mm
    line("Issued by: Quality Assurance Manager", "Helvetica", 10)
    c.showPage()
    c.save()


def write_expected(path: Path) -> None:
    """샘플별 기대 판정(문서·항목)을 expected.json으로 저장한다. 이후 테스트의 정답지."""
    data = {
        "_설명": "샘플 COA별 기대 판정. scripts/make_samples.py가 생성하며 테스트 정답지로 쓴다.",
        "spec_version": SPEC_VERSION,
        "samples": {
            s.file: {
                "supplier": s.supplier,
                "material_code": MATERIAL_CODE,
                "lot": s.lot,
                "form": s.form,
                "overall": s.overall,
                "items": {
                    i.std: {"raw_name": i.name, "raw_value": i.value, "raw_unit": i.unit,
                            "status": i.status, "note": i.note}
                    for i in s.items
                },
            }
            for s in SAMPLES
        },
    }
    path.write_text(json.dumps(data, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")


def build_all(samples_dir: Path, spec_path: Path) -> list[Path]:
    """규격 마스터, 샘플 PDF 5장, expected.json을 모두 만들고 PDF 경로 목록을 반환한다."""
    samples_dir.mkdir(parents=True, exist_ok=True)
    write_spec_master(spec_path)
    pdfs = []
    for s in SAMPLES:
        path = samples_dir / s.file
        (build_alpha_pdf if s.form == "alpha" else build_nippon_pdf)(path, s)
        pdfs.append(path)
    write_expected(samples_dir / "expected.json")
    return pdfs


def verify_pdfs(samples_dir: Path) -> bool:
    """pdfplumber로 각 PDF 텍스트를 추출해 Lot·항목값이 읽히는지 출력한다. 모두 읽히면 True."""
    ok = True
    for s in SAMPLES:
        with pdfplumber.open(samples_dir / s.file) as pdf:
            page = pdf.pages[0]
            text = page.extract_text() or ""
            tables = page.extract_tables()
        found = [i.name for i in s.items if f"{i.name}" in text and i.value in text]
        lot_ok = s.lot in text
        good = lot_ok and len(found) == len(s.items)
        ok &= good
        print(f"[{'확인' if good else '실패'}] {s.file}: 글자 {len(text)}자, 표 {len(tables)}개, "
              f"Lot {'읽힘' if lot_ok else '없음'}, 항목값 {len(found)}/{len(s.items)} 읽힘, 기대 판정 {s.overall}")
        for row in text.splitlines():
            if any(row.startswith(i.name) for i in s.items):
                print(f"      {row}")
    return ok


def main() -> int:
    """설정 경로에 샘플 데이터를 만들고 텍스트 추출을 확인한다."""
    paths = load_config()["paths"]
    samples_dir, spec_path = paths["samples"], paths["spec_master"]
    build_all(samples_dir, spec_path)
    print(f"[생성] {spec_path}")
    print(f"[생성] {samples_dir / 'expected.json'}")
    print(f"[생성] {samples_dir} 에 PDF {len(SAMPLES)}장")
    print("── pdfplumber 텍스트 추출 확인 ──")
    return 0 if verify_pdfs(samples_dir) else 1


if __name__ == "__main__":
    sys.exit(main())
